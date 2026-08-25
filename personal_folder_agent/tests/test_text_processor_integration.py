"""Integration tests for the Text Processor component.

Tests cover (Requirements 2.1, 2.2, 2.6):
- Plain text extraction with control-character removal and whitespace normalization
- Format-specific extraction for DOCX, XLSX, PPTX, CSV, and PDF
- Plain-text fallback for unsupported formats (success and failure paths)
- Missing-file handling (never raises)
- Supported-format registry contents

Each test creates a real sample file on disk (via tmp_path) and runs the
ConcreteTextProcessor against it, exercising the full extraction path rather
than mocking the format libraries.
"""

import os
from datetime import datetime

import pytest

from personal_folder_agent.components.text_processor import ConcreteTextProcessor
from personal_folder_agent.models import FileMetadata


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def make_metadata(path: str, extension: str) -> FileMetadata:
    """Build a FileMetadata instance for a file on disk.

    Args:
        path: Absolute path to the file.
        extension: File extension including the dot (e.g. '.txt').

    Returns:
        FileMetadata with size, extension, timestamp, hash, and permissions.
    """
    size_bytes = os.path.getsize(path)
    return FileMetadata(
        path=path,
        size_bytes=size_bytes,
        extension=extension,
        last_modified=datetime.now(),
        sha256_hash="0" * 64,
        permissions="rw-r--r--",
    )


def _minimal_pdf(text: str) -> bytes:
    """Build a minimal single-page PDF containing the given text.

    Hand-crafts a valid PDF 1.4 document with a single Helvetica text object
    so the PDFHandler (pdfplumber) has real content to extract. pdfplumber
    cannot create PDFs, so this raw-byte template is used instead.

    Args:
        text: The text to place in the PDF content stream.

    Returns:
        PDF document bytes.
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
    ]
    stream = ("BT /F1 24 Tf 100 700 Td (" + text + ") Tj ET").encode("latin-1")
    objects.append(
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"

    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += (
        b"trailer\n<< /Root 1 0 R /Size "
        + str(len(objects) + 1).encode()
        + b" >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


@pytest.fixture
def processor():
    """A fresh ConcreteTextProcessor for each test."""
    return ConcreteTextProcessor()


# ---------------------------------------------------------------------------
# 1. Plain text extraction (Requirement 2.1)
# ---------------------------------------------------------------------------


def test_txt_extraction_cleans_control_chars_and_whitespace(tmp_path, processor):
    """TXT extraction removes control chars and normalizes whitespace."""
    path = tmp_path / "sample.txt"
    # Include control chars (NUL, unit separator), tabs, CRLF, and messy spaces
    path.write_text(
        "Hello\x00World\t\twith\ttabs\r\n"
        "Second   line   with   extra   spaces\x1f\n"
        "  leading and trailing  \n",
        encoding="utf-8",
    )

    result = processor.extract_text(str(path), make_metadata(str(path), ".txt"))

    assert result.error_message is None
    assert result.extraction_method == "plaintext"
    assert "\x00" not in result.extracted_text
    assert "\x1f" not in result.extracted_text
    assert "\r" not in result.extracted_text
    # NUL is removed (not replaced), so "Hello" and "World" join directly
    assert "HelloWorld with tabs" in result.extracted_text
    # Tabs collapsed to single spaces; runs of spaces collapsed
    assert "Second line with extra spaces" in result.extracted_text
    assert "leading and trailing" in result.extracted_text
    # No leading/trailing whitespace on the whole result
    assert result.extracted_text == result.extracted_text.strip()


# ---------------------------------------------------------------------------
# 2. DOCX extraction (Requirement 2.2)
# ---------------------------------------------------------------------------


def test_docx_extraction(tmp_path, processor):
    """DOCX extraction returns paragraph text via python-docx."""
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("First paragraph of the document.")
    doc.add_paragraph("Second paragraph with more content.")
    doc.add_paragraph("")  # empty paragraph should be skipped
    doc.save(str(path))

    result = processor.extract_text(str(path), make_metadata(str(path), ".docx"))

    assert result.error_message is None
    assert result.extraction_method == "python-docx"
    assert "First paragraph of the document." in result.extracted_text
    assert "Second paragraph with more content." in result.extracted_text


# ---------------------------------------------------------------------------
# 3. XLSX extraction (Requirement 2.2)
# ---------------------------------------------------------------------------


def test_xlsx_extraction(tmp_path, processor):
    """XLSX extraction returns cell values and sheet headers."""
    from openpyxl import Workbook

    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "Alpha"
    ws["B2"] = 42
    wb.save(str(path))

    result = processor.extract_text(str(path), make_metadata(str(path), ".xlsx"))

    assert result.error_message is None
    assert result.extraction_method == "openpyxl"
    assert "[Sheet: Data]" in result.extracted_text
    assert "Name" in result.extracted_text
    assert "Value" in result.extracted_text
    assert "Alpha" in result.extracted_text
    assert "42" in result.extracted_text


# ---------------------------------------------------------------------------
# 4. PPTX extraction (Requirement 2.2)
# ---------------------------------------------------------------------------


def test_pptx_extraction(tmp_path, processor):
    """PPTX extraction returns slide text via python-pptx."""
    from pptx import Presentation

    path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Presentation Title"
    body = slide.placeholders[1]
    body.text = "Body text for the first slide."
    prs.save(str(path))

    result = processor.extract_text(str(path), make_metadata(str(path), ".pptx"))

    assert result.error_message is None
    assert result.extraction_method == "python-pptx"
    assert "[Slide 1]" in result.extracted_text
    assert "Presentation Title" in result.extracted_text
    assert "Body text for the first slide." in result.extracted_text


# ---------------------------------------------------------------------------
# 5. CSV extraction (Requirement 2.2)
# ---------------------------------------------------------------------------


def test_csv_extraction(tmp_path, processor):
    """CSV extraction returns row values via pandas."""
    path = tmp_path / "sample.csv"
    path.write_text(
        "name,age,city\nAlice,30,New York\nBob,25,London\n",
        encoding="utf-8",
    )

    result = processor.extract_text(str(path), make_metadata(str(path), ".csv"))

    assert result.error_message is None
    assert result.extraction_method == "pandas_csv"
    assert "Alice" in result.extracted_text
    assert "Bob" in result.extracted_text
    assert "New York" in result.extracted_text
    assert "London" in result.extracted_text


# ---------------------------------------------------------------------------
# 6. PDF extraction (Requirement 2.2)
# ---------------------------------------------------------------------------


def test_pdf_extraction(tmp_path, processor):
    """PDF extraction returns text via pdfplumber from a minimal PDF."""
    path = tmp_path / "sample.pdf"
    path.write_bytes(_minimal_pdf("Hello PDF"))

    result = processor.extract_text(str(path), make_metadata(str(path), ".pdf"))

    assert result.error_message is None
    assert result.extraction_method == "pdfplumber"
    assert "Hello PDF" in result.extracted_text


# ---------------------------------------------------------------------------
# 7. Fallback for unsupported formats (Requirement 2.6)
# ---------------------------------------------------------------------------


def test_unsupported_text_extension_falls_back_to_plaintext(tmp_path, processor):
    """Unsupported extension with plain-text content extracts via fallback."""
    path = tmp_path / "notes.xyz"
    path.write_text("This is plain text in an unknown format.", encoding="utf-8")

    result = processor.extract_text(str(path), make_metadata(str(path), ".xyz"))

    assert result.error_message is None
    assert result.extraction_method == "plaintext_fallback"
    assert "This is plain text in an unknown format." in result.extracted_text


def test_unsupported_binary_extension_fallback_fails(tmp_path, processor):
    """Unsupported binary extension sets an error and empty extracted text."""
    path = tmp_path / "blob.bin"
    path.write_bytes(os.urandom(256))

    result = processor.extract_text(str(path), make_metadata(str(path), ".bin"))

    assert result.extracted_text == ""
    assert result.error_message is not None
    assert result.error_message.startswith("Plain text fallback extraction failed")


# ---------------------------------------------------------------------------
# 8. Missing file handling
# ---------------------------------------------------------------------------


def test_missing_file_returns_error_result(tmp_path, processor):
    """Extracting a nonexistent path returns an error result, never raises."""
    missing = tmp_path / "does_not_exist.txt"
    metadata = FileMetadata(
        path=str(missing),
        size_bytes=0,
        extension=".txt",
        last_modified=datetime.now(),
        sha256_hash="0" * 64,
        permissions="rw-r--r--",
    )

    result = processor.extract_text(str(missing), metadata)

    assert result.error_message is not None
    assert result.extracted_text == ""


# ---------------------------------------------------------------------------
# 9. Supported formats registry
# ---------------------------------------------------------------------------


def test_get_supported_formats_includes_expected_extensions(processor):
    """The registry reports the expected set of supported extensions."""
    formats = processor.get_supported_formats()

    assert formats == sorted(formats)
    for ext in (".pdf", ".docx", ".xlsx", ".pptx", ".csv", ".txt"):
        assert ext in formats
